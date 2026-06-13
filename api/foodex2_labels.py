"""
IDRISK2 — Resolução de rótulos FoodEx2.

Carrega, a partir do XLSX da taxonomia (Appendix B):
- as dimensões das facetas (sheet 'attribute'): F01 -> "Source", F04 -> "Ingredient", …
- os nomes dos termos (sheet 'term'): termCode -> termExtendedName (A0F4B -> "Flagfish").

Permite mostrar facetas legíveis no frontend em vez de códigos crus.
Carregamento preguiçoso (lazy) e em cache; thread-safe.
"""
from __future__ import annotations

import glob
import logging
import os
import threading
from pathlib import Path
from typing import Any, Optional

logger = logging.getLogger(__name__)

_lock = threading.Lock()
_facet_dims: Optional[dict[str, str]] = None   # F-code -> rótulo da dimensão
_term_labels: Optional[dict[str, str]] = None  # termCode -> nome


def _xlsx_path() -> Optional[str]:
    docs = os.environ.get("IDRISK2_DOCS_DIR", "./docs/efsa")
    matches = glob.glob(str(Path(docs) / "*appendix*.xlsx"))
    return matches[0] if matches else None


def ensure_loaded() -> None:
    global _facet_dims, _term_labels
    if _term_labels is not None:
        return
    with _lock:
        if _term_labels is not None:
            return
        dims: dict[str, str] = {}
        terms: dict[str, str] = {}
        path = _xlsx_path()
        if path:
            try:
                import openpyxl

                wb = openpyxl.load_workbook(path, read_only=True)
                # Dimensões das facetas
                ws = wb["attribute"]
                hdr = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
                ci = {h: i for i, h in enumerate(hdr)}
                for row in ws.iter_rows(min_row=2, values_only=True):
                    code = row[ci["code"]]
                    label = row[ci["label"]] or row[ci["name"]]
                    if code:
                        dims[str(code)] = str(label) if label else str(code)
                # Nomes dos termos
                ws2 = wb["term"]
                hdr2 = list(next(ws2.iter_rows(min_row=1, max_row=1, values_only=True)))
                ti = {h: i for i, h in enumerate(hdr2)}
                for row in ws2.iter_rows(min_row=2, values_only=True):
                    code = row[ti["termCode"]]
                    name = row[ti["termExtendedName"]]
                    if code:
                        terms[str(code)] = str(name) if name else ""
                wb.close()
                logger.info(
                    "FoodEx2 labels carregados: %d dimensões, %d termos.",
                    len(dims),
                    len(terms),
                )
            except Exception as exc:  # noqa: BLE001
                logger.warning("Não foi possível carregar rótulos FoodEx2: %s", exc)
        _facet_dims = dims
        _term_labels = terms


def facet_group_label(code: str) -> str:
    ensure_loaded()
    return (_facet_dims or {}).get(code, code)


def term_label(code: str) -> Optional[str]:
    ensure_loaded()
    return (_term_labels or {}).get(code)


def resolve_facets(facets: dict[str, str] | None) -> list[dict[str, Any]]:
    """Converte {F01: A0F4B, ...} em [{group, group_label, code, label}, ...]."""
    ensure_loaded()
    out: list[dict[str, Any]] = []
    for group, code in (facets or {}).items():
        out.append(
            {
                "group": group,
                "group_label": (_facet_dims or {}).get(group, group),
                "code": code,
                "label": (_term_labels or {}).get(code),
            }
        )
    return out
